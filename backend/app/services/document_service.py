import asyncio
import os
import re
from typing import List, Dict, Any
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, delete
from app.models.models import Document, DocumentPage, DocumentChunk, WorkbookExercise

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None


class DocumentService:
    @staticmethod
    async def process_document(db: AsyncSession, course_id: int, document_type: str, filename: str, file_path: str, file_data: bytes = None) -> Document:
        file_ext = os.path.splitext(filename)[1].lower()
        pages_content = []

        if file_ext == '.pdf':
            pages_content = DocumentService._extract_pdf(file_path, file_data)
        elif file_ext in ['.docx', '.doc']:
            pages_content = DocumentService._extract_docx(file_path, file_data)
        elif file_ext == '.txt':
            pages_content = DocumentService._extract_txt(file_path, file_data)
        else:
            raise ValueError(f"Unsupported file type: {file_ext}")

        doc = Document(
            course_id=course_id,
            document_type=document_type,
            filename=filename,
            file_path=file_path,
            file_data=file_data,
            total_pages=len(pages_content),
            status="processed"
        )
        db.add(doc)
        await db.flush()

        for page_num, text in enumerate(pages_content, start=1):
            doc_page = DocumentPage(
                document_id=doc.id,
                page_number=page_num,
                content=text,
                clean_text=text.strip()
            )
            db.add(doc_page)

            chunks = DocumentService._parse_page_chunks(doc.id, page_num, text, document_type)
            for chunk in chunks:
                db.add(chunk)

            if document_type.lower() == 'workbook':
                exercises = DocumentService._parse_workbook_exercises(course_id, doc.id, page_num, text)
                for ex in exercises:
                    db.add(ex)

        await db.commit()
        await db.refresh(doc)
        return doc

    @staticmethod
    async def process_uploaded_document(db: AsyncSession, doc: Document) -> None:
        """Parse an already-stored (pending) document and build its search rows.

        The CPU-bound extraction runs in a worker thread so the event loop (and
        therefore the whole API on a single-instance server) is not blocked on
        big PDFs. Existing page/chunk/exercise rows are replaced so re-running
        is safe and idempotent.
        """
        file_ext = os.path.splitext(doc.filename)[1].lower()
        pages_content = await asyncio.to_thread(
            DocumentService._extract_for_ext,
            file_ext,
            doc.file_data or b"",
            doc.filename,
        )

        await db.execute(delete(DocumentPage).where(DocumentPage.document_id == doc.id))
        await db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc.id))
        await db.execute(delete(WorkbookExercise).where(WorkbookExercise.document_id == doc.id))

        for page_num, text in enumerate(pages_content, start=1):
            db.add(
                DocumentPage(
                    document_id=doc.id,
                    page_number=page_num,
                    content=text,
                    clean_text=text.strip(),
                )
            )
            for chunk in DocumentService._parse_page_chunks(doc.id, page_num, text, doc.document_type):
                db.add(chunk)
            if doc.document_type.lower() == "workbook":
                for ex in DocumentService._parse_workbook_exercises(doc.course_id, doc.id, page_num, text):
                    db.add(ex)

        doc.total_pages = len(pages_content)
        doc.status = "processed"

    @staticmethod
    def _extract_for_ext(file_ext: str, file_data: bytes, filename: str) -> List[str]:
        if file_data and file_ext == ".pdf" and not file_data.startswith(b"%PDF"):
            try:
                from app.services.upload_codec import decompress_payload
                file_data = decompress_payload(file_data, compressed=True)
            except Exception:
                pass

        if file_ext == ".pdf":
            return DocumentService._extract_pdf("", file_data)
        if file_ext in (".docx", ".doc"):
            return DocumentService._extract_docx("", file_data)
        if file_ext == ".txt":
            return DocumentService._extract_txt("", file_data)
        raise ValueError(f"Unsupported file type: {file_ext}")

    @staticmethod
    def _extract_pdf(file_path: str, file_data: bytes = None) -> List[str]:
        pages = []
        if file_data and not file_data.startswith(b"%PDF"):
            try:
                from app.services.upload_codec import decompress_payload
                file_data = decompress_payload(file_data, compressed=True)
            except Exception:
                pass

        if pypdf:
            import io
            stream = io.BytesIO(file_data) if file_data else file_path
            try:
                reader = pypdf.PdfReader(stream, strict=False)
                for page in reader.pages:
                    try:
                        text = page.extract_text() or ""
                    except Exception:
                        text = ""
                    pages.append(text)
            except Exception:
                if file_data:
                    pages = [file_data.decode('utf-8', errors='ignore')]
                else:
                    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                        pages = [f.read()]
        else:
            if file_data:
                pages = [file_data.decode('utf-8', errors='ignore')]
            else:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    pages = [f.read()]
        return pages

    @staticmethod
    def _extract_docx(file_path: str, file_data: bytes = None) -> List[str]:
        pages = []
        if docx:
            if file_data:
                import io
                doc = docx.Document(io.BytesIO(file_data))
            else:
                doc = docx.Document(file_path)
            current_page_text = []
            for p in doc.paragraphs:
                current_page_text.append(p.text)
                if len('\n'.join(current_page_text).split()) > 400:
                    pages.append('\n'.join(current_page_text))
                    current_page_text = []
            if current_page_text:
                pages.append('\n'.join(current_page_text))
        else:
            if file_data:
                pages = [file_data.decode('utf-8', errors='ignore')]
            else:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    pages = [f.read()]
        return pages if pages else [""]

    @staticmethod
    def _extract_txt(file_path: str, file_data: bytes = None) -> List[str]:
        if file_data:
            content = file_data.decode('utf-8', errors='ignore')
        else:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
        chunk_size = 3000
        pages = [content[i:i+chunk_size] for i in range(0, len(content), chunk_size)]
        return pages if pages else [""]

    @staticmethod
    def _parse_page_chunks(doc_id: int, page_num: int, text: str, doc_type: str) -> List[DocumentChunk]:
        chunks = []
        unit_match = re.search(r'Unit\s*(\d+|[I|V|X]+)', text, re.IGNORECASE)
        chapter_match = re.search(r'Chapter\s*(\d+|[I|V|X]+)', text, re.IGNORECASE)
        lesson_match = re.search(r'Lesson\s*(\d+|[I|V|X]+)', text, re.IGNORECASE)
        section_match = re.search(r'(Vocabulary|Grammar|Reading|Writing|Activities|Practice|Exercise)', text, re.IGNORECASE)

        unit = f"Unit {unit_match.group(1)}" if unit_match else None
        chapter = f"Chapter {chapter_match.group(1)}" if chapter_match else None
        lesson = f"Lesson {lesson_match.group(1)}" if lesson_match else None
        section = section_match.group(1) if section_match else None

        chunk = DocumentChunk(
            document_id=doc_id,
            page_number=page_num,
            unit=unit,
            chapter=chapter,
            lesson=lesson,
            section=section,
            content=text,
            metadata_json={
                "doc_type": doc_type,
                "has_vocabulary": "vocabulary" in text.lower(),
                "has_grammar": "grammar" in text.lower(),
                "has_reading": "reading" in text.lower(),
                "has_exercises": "exercise" in text.lower() or "activity" in text.lower()
            }
        )
        chunks.append(chunk)
        return chunks

    @staticmethod
    def _parse_workbook_exercises(course_id: int, doc_id: int, page_num: int, text: str) -> List[WorkbookExercise]:
        exercises = []
        matches = re.finditer(r'(Exercise|Activity|Task)\s*(\d+[a-z]?)\s*[:\.]?\s*([^\n]+)', text, re.IGNORECASE)
        unit_match = re.search(r'Unit\s*(\d+|[I|V|X]+)', text, re.IGNORECASE)
        unit_str = f"Unit {unit_match.group(1)}" if unit_match else None

        for match in matches:
            ex_type = match.group(1)
            ex_num = match.group(2)
            prompt = match.group(3).strip()
            wb_ex = WorkbookExercise(
                course_id=course_id,
                document_id=doc_id,
                page_number=page_num,
                unit=unit_str,
                exercise_number=f"{ex_type} {ex_num}",
                title=f"{ex_type} {ex_num}",
                prompt=prompt,
                metadata_json={"page": page_num}
            )
            exercises.append(wb_ex)
        return exercises

    @staticmethod
    async def search_documents(db: AsyncSession, course_id: int, query: str, document_type: str = None) -> List[Dict[str, Any]]:
        # Structured content search across document chunks
        stmt = (
            select(DocumentChunk, Document)
            .join(Document, DocumentChunk.document_id == Document.id)
            .where(Document.course_id == course_id)
        )
        if document_type:
            stmt = stmt.where(Document.document_type == document_type)

        result = await db.execute(stmt)
        rows = result.all()

        query_terms = query.lower().split()
        matched_chunks = []

        for chunk, doc in rows:
            content_lower = chunk.content.lower()
            score = sum(1 for term in query_terms if term in content_lower)
            if score > 0 or not query:
                matched_chunks.append({
                    "document_name": doc.filename,
                    "document_type": doc.document_type,
                    "page_number": chunk.page_number,
                    "unit": chunk.unit,
                    "chapter": chunk.chapter,
                    "lesson": chunk.lesson,
                    "section": chunk.section,
                    "content": chunk.content,
                    "relevance_score": score
                })

        matched_chunks.sort(key=lambda x: x["relevance_score"], reverse=True)
        return matched_chunks[:15]
